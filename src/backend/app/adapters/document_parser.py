import csv
import json
import os
from io import BytesIO
import re
from zipfile import ZipFile

from docx import Document as DocxDocument
from pypdf import PdfReader


class UnsupportedDocumentTypeError(ValueError):
    pass


class PlainTextDocumentParser:
    parser_name = "plain_text"

    def parse_content(self, content: str) -> list[str]:
        normalized = content.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not normalized:
            return []
        return [normalized]

    def parse_bytes(self, content: bytes, filename: str | None = None) -> list[str]:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ValueError("Text file must be UTF-8 encoded.") from exc
        return self.parse_content(text)


class MarkdownDocumentParser(PlainTextDocumentParser):
    parser_name = "markdown"

    def parse_content(self, content: str) -> list[str]:
        normalized = content.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not normalized:
            return []
        without_fences = re.sub(r"```[\s\S]*?```", lambda match: match.group(0).strip(), normalized)
        sections = [section.strip() for section in re.split(r"\n(?=#{1,6}\s+)", without_fences) if section.strip()]
        return sections or [normalized]


class PdfDocumentParser(PlainTextDocumentParser):
    parser_name = "pdf"

    def parse_bytes(self, content: bytes, filename: str | None = None) -> list[str]:
        reader = PdfReader(BytesIO(content))
        sections: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(f"[Page {index}]\n{text}")
        return sections or self._ocr_scanned_pdf(content)

    def _ocr_scanned_pdf(self, content: bytes) -> list[str]:
        try:
            import pytesseract
            import pypdfium2
        except ImportError as exc:
            raise ValueError(
                "OCR dependencies are unavailable. Install pytesseract, pypdfium2 and Pillow before uploading scanned PDFs."
            ) from exc

        tesseract_cmd = os.getenv("TESSERACT_CMD", "").strip()
        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

        language = os.getenv("TESSERACT_LANG", "chi_sim+eng").strip() or "chi_sim+eng"
        try:
            pdf = pypdfium2.PdfDocument(content)
            sections: list[str] = []
            for index in range(len(pdf)):
                page = pdf[index]
                image = page.render(scale=300 / 72).to_pil()
                text = pytesseract.image_to_string(image, lang=language).strip()
                if text:
                    sections.append(f"[Page {index + 1}]\n{text}")
        except pytesseract.TesseractNotFoundError as exc:
            raise ValueError(
                "OCR requires Tesseract. Install Tesseract and configure TESSERACT_CMD if it is not on PATH."
            ) from exc
        except pytesseract.TesseractError as exc:
            raise ValueError(
                f"OCR failed. Verify the Tesseract language data for '{language}' is installed and TESSERACT_LANG is correct."
            ) from exc
        except Exception as exc:
            raise ValueError("Unable to render the scanned PDF for OCR.") from exc

        if not sections:
            raise ValueError("OCR completed but no text was recognized in the scanned PDF.")
        return sections


class WordDocumentParser(PlainTextDocumentParser):
    parser_name = "docx"

    def parse_bytes(self, content: bytes, filename: str | None = None) -> list[str]:
        document = DocxDocument(BytesIO(content))
        sections: list[str] = []
        current: list[str] = []

        def flush() -> None:
            if current:
                text = "\n".join(item for item in current if item.strip()).strip()
                if text:
                    sections.append(text)
                current.clear()

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""
            if style.startswith("heading") and current:
                flush()
            current.append(text)
        flush()

        for table in document.tables:
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                sections.append("\n".join(rows))

        if not sections:
            raise ValueError("No extractable text found in the Word document.")
        return sections


class PowerPointDocumentParser(PlainTextDocumentParser):
    parser_name = "pptx"

    def parse_bytes(self, content: bytes, filename: str | None = None) -> list[str]:
        sections: list[str] = []
        with ZipFile(BytesIO(content)) as archive:
            slide_names = sorted(
                (name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")),
                key=_powerpoint_slide_sort_key,
            )
            for index, name in enumerate(slide_names, start=1):
                texts = _office_xml_texts(archive.read(name))
                if texts:
                    sections.append(f"[Slide {index}]\n" + "\n".join(texts))
        if not sections:
            raise ValueError("No extractable text found in the PowerPoint document.")
        return sections


class SpreadsheetDocumentParser(PlainTextDocumentParser):
    parser_name = "spreadsheet"

    def parse_bytes(self, content: bytes, filename: str | None = None) -> list[str]:
        suffix = (filename or "").rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
        if suffix == "csv":
            return self._parse_csv(content)
        return self._parse_xlsx(content)

    def _parse_csv(self, content: bytes) -> list[str]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV file must be UTF-8 encoded.") from exc
        rows = [" | ".join(cell.strip() for cell in row if cell.strip()) for row in csv.reader(text.splitlines())]
        rows = [row for row in rows if row]
        if not rows:
            raise ValueError("No extractable rows found in the CSV file.")
        return ["\n".join(rows)]

    def _parse_xlsx(self, content: bytes) -> list[str]:
        sections: list[str] = []
        with ZipFile(BytesIO(content)) as archive:
            shared_strings = self._shared_strings(archive)
            sheet_names = sorted(name for name in archive.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"))
            for index, name in enumerate(sheet_names, start=1):
                rows = self._sheet_rows(archive.read(name), shared_strings)
                if rows:
                    sections.append(f"[Sheet {index}]\n" + "\n".join(rows))
        if not sections:
            raise ValueError("No extractable rows found in the spreadsheet.")
        return sections

    def _shared_strings(self, archive: ZipFile) -> list[str]:
        if "xl/sharedStrings.xml" not in archive.namelist():
            return []
        root = _xml_root(archive.read("xl/sharedStrings.xml"))
        values: list[str] = []
        for item in root.iter():
            if item.tag.endswith("}si") or item.tag == "si":
                text = " ".join(child.text.strip() for child in item.iter() if child.tag.endswith("}t") and child.text and child.text.strip())
                values.append(text)
        return values

    def _sheet_rows(self, payload: bytes, shared_strings: list[str]) -> list[str]:
        root = _xml_root(payload)
        rows: list[str] = []
        for row in root.iter():
            if not (row.tag.endswith("}row") or row.tag == "row"):
                continue
            cells: list[str] = []
            for cell in row:
                if not (cell.tag.endswith("}c") or cell.tag == "c"):
                    continue
                cell_type = cell.attrib.get("t")
                value = self._cell_value(cell, cell_type, shared_strings)
                if value:
                    cells.append(value)
            if cells:
                rows.append(" | ".join(cells))
        return rows

    def _cell_value(self, cell, cell_type: str | None, shared_strings: list[str]) -> str:
        if cell_type == "inlineStr":
            return " ".join(item.text.strip() for item in cell.iter() if item.tag.endswith("}t") and item.text and item.text.strip())
        value_node = next((item for item in cell if item.tag.endswith("}v") or item.tag == "v"), None)
        if value_node is None or value_node.text is None:
            return ""
        raw = value_node.text.strip()
        if cell_type == "s" and raw.isdigit():
            index = int(raw)
            return shared_strings[index] if 0 <= index < len(shared_strings) else raw
        return raw


class NotebookDocumentParser(PlainTextDocumentParser):
    parser_name = "notebook"

    def parse_bytes(self, content: bytes, filename: str | None = None) -> list[str]:
        try:
            notebook = json.loads(content.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("Notebook file must be a valid UTF-8 .ipynb JSON document.") from exc
        sections: list[str] = []
        for index, cell in enumerate(notebook.get("cells", []), start=1):
            source = cell.get("source", "")
            text = "".join(source) if isinstance(source, list) else str(source)
            text = text.strip()
            if text:
                cell_type = cell.get("cell_type", "cell")
                sections.append(f"[{cell_type} cell {index}]\n{text}")
        if not sections:
            raise ValueError("No extractable cells found in the notebook.")
        return sections


class CodeDocumentParser(PlainTextDocumentParser):
    parser_name = "code_text"

    def parse_bytes(self, content: bytes, filename: str | None = None) -> list[str]:
        sections = super().parse_bytes(content, filename)
        if filename and sections:
            return [f"[Code file: {filename}]\n{section}" for section in sections]
        return sections


class SubtitleDocumentParser(PlainTextDocumentParser):
    parser_name = "subtitle"


def _xml_root(payload: bytes):
    from xml.etree import ElementTree

    return ElementTree.fromstring(payload)


def _office_xml_texts(payload: bytes) -> list[str]:
    root = _xml_root(payload)
    return [item.text.strip() for item in root.iter() if item.tag.endswith("}t") and item.text and item.text.strip()]


def _powerpoint_slide_sort_key(name: str) -> int:
    match = re.search(r"slide(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


CODE_SUFFIXES = {
    "py",
    "java",
    "js",
    "jsx",
    "ts",
    "tsx",
    "html",
    "css",
    "json",
    "yaml",
    "yml",
    "sql",
    "c",
    "cpp",
    "h",
    "hpp",
    "go",
    "rs",
}


def build_document_parser(source_type: str) -> PlainTextDocumentParser:
    normalized = source_type.strip().lower()
    if normalized in {"markdown", "md", "text/markdown"}:
        return MarkdownDocumentParser()
    if normalized in {"text", "txt", "plain_text", "text/plain"}:
        return PlainTextDocumentParser()
    if normalized in {"pdf", "application/pdf"}:
        return PdfDocumentParser()
    if normalized in {
        "docx",
        "word",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }:
        return WordDocumentParser()
    if normalized in {"pptx", "powerpoint", "application/vnd.openxmlformats-officedocument.presentationml.presentation"}:
        return PowerPointDocumentParser()
    if normalized in {"xlsx", "spreadsheet", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}:
        return SpreadsheetDocumentParser()
    if normalized in {"csv", "text/csv"}:
        return SpreadsheetDocumentParser()
    if normalized in {"ipynb", "notebook", "application/x-ipynb+json"}:
        return NotebookDocumentParser()
    if normalized in {"srt", "vtt", "text/vtt", "subtitle"}:
        return SubtitleDocumentParser()
    if normalized in CODE_SUFFIXES or normalized.startswith("text/x-"):
        return CodeDocumentParser()
    raise UnsupportedDocumentTypeError(f"Unsupported document source_type: {source_type}")


def infer_source_type(filename: str, content_type: str | None = None) -> str:
    normalized_type = (content_type or "").strip().lower()
    if normalized_type in {
        "text/plain",
        "text/markdown",
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/csv",
        "text/vtt",
    }:
        return normalized_type

    suffix = filename.rsplit(".", 1)[-1].strip().lower() if "." in filename else ""
    if suffix == "ppt" or normalized_type == "application/vnd.ms-powerpoint":
        raise UnsupportedDocumentTypeError("Legacy .ppt files are not supported. Convert the presentation to .pptx before uploading.")
    if suffix in {"txt", "text"}:
        return "text"
    if suffix in {"md", "markdown"}:
        return "markdown"
    if suffix == "pdf":
        return "pdf"
    if suffix == "docx":
        return "docx"
    if suffix == "pptx":
        return "pptx"
    if suffix == "xlsx":
        return "xlsx"
    if suffix == "csv":
        return "csv"
    if suffix == "ipynb":
        return "ipynb"
    if suffix in {"srt", "vtt"}:
        return suffix
    if suffix in CODE_SUFFIXES:
        return suffix
    raise UnsupportedDocumentTypeError(f"Unsupported document file type: {filename}")
